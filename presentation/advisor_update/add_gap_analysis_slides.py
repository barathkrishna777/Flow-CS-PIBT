"""Append gap-analysis slides to the April 24 Flow-CS-PIBT advisor deck.

This script intentionally creates and modifies a copy of the input deck. The
original presentation is never overwritten.

Run from the repository root with an environment that has python-pptx:

    python presentation/advisor_update/add_gap_analysis_slides.py

The generated deck is written to:

    presentation/advisor_update/outputs/flow_cs_pibt_barath_2026-04-24_gap_analysis.pptx
"""

from __future__ import annotations

import shutil
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.util import Inches, Pt
except ImportError as exc:  # pragma: no cover - user-facing dependency guard
    raise SystemExit(
        "Missing dependency: python-pptx. Install it with:\n"
        "    python -m pip install python-pptx\n"
    ) from exc


REPO_ROOT = Path(__file__).resolve().parents[2]
DECK_DIR = REPO_ROOT / "presentation" / "advisor_update" / "outputs"
INPUT_PPTX = DECK_DIR / "flow_cs_pibt_barath_2026-04-24.pptx"
OUTPUT_PPTX = DECK_DIR / "flow_cs_pibt_barath_2026-04-24_gap_analysis.pptx"


# CMU-inspired palette used by the source deck.
CMU_RED = RGBColor(196, 18, 48)
DARK = RGBColor(36, 47, 58)
GRAY = RGBColor(91, 103, 112)
LIGHT_GRAY = RGBColor(241, 243, 245)
MID_GRAY = RGBColor(194, 202, 208)
BLUE = RGBColor(0, 59, 113)
GOLD = RGBColor(255, 184, 28)
GREEN = RGBColor(0, 132, 61)
TEAL = RGBColor(0, 153, 153)
WHITE = RGBColor(255, 255, 255)

FONT = "Arial"


def add_text(
    slide,
    text: str,
    left,
    top,
    width,
    height,
    *,
    size: int = 18,
    color: RGBColor = DARK,
    bold: bool = False,
    italic: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
):
    """Add a styled text box and return the shape."""
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def add_title(slide, title: str, subtitle: str | None = None):
    """Match the source deck's title zone: dark title, gray subtitle, red rule."""
    title_size = 21 if len(title) <= 54 else 18
    add_text(slide, title, Inches(0.42), Inches(0.34), Inches(9.0), Inches(0.42), size=title_size, bold=True)
    if subtitle:
        add_text(slide, subtitle, Inches(0.42), Inches(0.84), Inches(9.0), Inches(0.28), size=11, color=GRAY)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.42), Inches(1.23), Inches(9.15), Inches(0.02))
    rule.fill.solid()
    rule.fill.fore_color.rgb = CMU_RED
    rule.line.fill.background()


def add_custom_legend(slide, items: list[tuple[str, RGBColor]], left, top, *, columns: int = 1):
    """Add a compact, editable legend that does not consume chart plot area."""
    for i, (label, color) in enumerate(items):
        col = i % columns
        row = i // columns
        x = left + Inches(1.92 * col)
        y = top + Inches(0.22 * row)
        swatch = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y + Inches(0.05), Inches(0.12), Inches(0.08))
        swatch.fill.solid()
        swatch.fill.fore_color.rgb = color
        swatch.line.fill.background()
        add_text(slide, label, x + Inches(0.16), y, Inches(1.65), Inches(0.2), size=8, color=GRAY)


def add_label(slide, text: str, left, top, width, height, color: RGBColor = CMU_RED):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = color
    shape.line.width = Pt(1.1)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.06)
    tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = color
    return shape


def add_callout(slide, heading: str, body: str, left, top, width, height, accent: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(253, 253, 253)
    shape.line.color.rgb = accent
    shape.line.width = Pt(1.1)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.16)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.08)
    p1 = tf.paragraphs[0]
    r1 = p1.add_run()
    r1.text = heading
    r1.font.name = FONT
    r1.font.size = Pt(14)
    r1.font.bold = True
    r1.font.color.rgb = accent
    p2 = tf.add_paragraph()
    p2.space_before = Pt(6)
    r2 = p2.add_run()
    r2.text = body
    r2.font.name = FONT
    r2.font.size = Pt(12)
    r2.font.color.rgb = DARK
    return shape


def add_placeholder(slide, label: str, left, top, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT_GRAY
    shape.fill.transparency = 20
    shape.line.color.rgb = MID_GRAY
    shape.line.width = Pt(1.4)
    shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = FONT
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = GRAY
    return shape


def add_clustered_column_chart(
    slide,
    categories: list[str],
    series: list[tuple[str, list[float], RGBColor]],
    left,
    top,
    width,
    height,
    *,
    value_max: int = 100,
    legend: bool = True,
):
    chart_data = CategoryChartData()
    chart_data.categories = categories
    for name, values, _color in series:
        chart_data.add_series(name, values)

    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, left, top, width, height, chart_data
    )
    chart = graphic_frame.chart
    chart.has_title = False
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False

    value_axis = chart.value_axis
    value_axis.minimum_scale = 0
    value_axis.maximum_scale = value_max
    value_axis.major_unit = 20
    value_axis.has_major_gridlines = True
    value_axis.tick_labels.font.size = Pt(9)
    value_axis.tick_labels.font.color.rgb = GRAY

    category_axis = chart.category_axis
    category_axis.tick_labels.font.size = Pt(9)
    category_axis.tick_labels.font.color.rgb = DARK

    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.number_format = '0"%"'
    plot.data_labels.font.size = Pt(9)
    plot.gap_width = 70

    for chart_series, (_, _values, color) in zip(chart.series, series):
        chart_series.format.fill.solid()
        chart_series.format.fill.fore_color.rgb = color
        chart_series.format.line.color.rgb = color
    return graphic_frame


def add_horizontal_arrow(slide, left, top, width, color: RGBColor = GRAY):
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, Inches(0.22))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = color
    arrow.line.fill.background()
    return arrow


def slide_gap_headline(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    add_title(
        slide,
        "Gap analysis: SSIL still wins the coordination cases",
        "The average gap is modest, but the hard-map failure mode is concentrated.",
    )
    add_clustered_column_chart(
        slide,
        ["Full eval", "Medium eval"],
        [
            ("Flow-vector policy", [64.05, 64.59], CMU_RED),
            ("Discrete imitation (SSIL)", [66.43, 67.84], BLUE),
        ],
        Inches(0.55),
        Inches(1.72),
        Inches(5.75),
        Inches(2.95),
        legend=False,
    )
    add_custom_legend(
        slide,
        [("Flow-vector policy", CMU_RED), ("Discrete imitation (SSIL)", BLUE)],
        Inches(0.72),
        Inches(1.43),
        columns=2,
    )
    add_callout(
        slide,
        "Headline target",
        "Beat 66.43% on the same Rishi-style full evaluation.",
        Inches(6.58),
        Inches(1.65),
        Inches(2.65),
        Inches(0.95),
        CMU_RED,
    )
    add_callout(
        slide,
        "Truth signal",
        "Medium eval tracks real progress better than quick sweeps.",
        Inches(6.58),
        Inches(2.82),
        Inches(2.65),
        Inches(0.95),
        BLUE,
    )
    add_text(
        slide,
        "Use quick runs only to filter; promote candidates based on medium-map behavior.",
        Inches(6.58),
        Inches(4.05),
        Inches(2.65),
        Inches(0.55),
        size=12,
        color=GRAY,
    )


def slide_dense_coordination_gap(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    add_title(
        slide,
        "The gap localizes to dense coordination",
        "random-32-32-10 is the clearest stress test for waiting, yielding, and bottlenecks.",
    )
    add_clustered_column_chart(
        slide,
        ["random-32-32-10"],
        [
            ("Flow-vector baseline", [50], CMU_RED),
            ("Discrete imitation (SSIL)", [90], BLUE),
            ("Hybrid discrete-primary", [75], GREEN),
        ],
        Inches(0.58),
        Inches(1.82),
        Inches(5.25),
        Inches(2.72),
        legend=False,
    )
    add_custom_legend(
        slide,
        [
            ("Flow-vector", CMU_RED),
            ("SSIL", BLUE),
            ("Hybrid", GREEN),
        ],
        Inches(0.78),
        Inches(1.42),
        columns=3,
    )
    add_placeholder(
        slide,
        "Plot placeholder:\ncongestion heatmap or failure rollout strip",
        Inches(6.12),
        Inches(1.62),
        Inches(3.02),
        Inches(1.65),
    )
    add_callout(
        slide,
        "What changed",
        "random-32 improves from 50% to 75%.",
        Inches(6.12),
        Inches(3.55),
        Inches(3.02),
        Inches(0.82),
        GREEN,
    )
    add_text(
        slide,
        "The problem is not only path quality; it is local coordination under crowd pressure.",
        Inches(0.7),
        Inches(4.62),
        Inches(8.4),
        Inches(0.38),
        size=13,
        color=DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def slide_training_objective(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    add_title(
        slide,
        "Working hypothesis: the trunk learns what the loss asks for",
        "SSIL trains discrete coordination from epoch 0; flow matching learns smooth direction first.",
    )
    y = Inches(1.65)
    add_callout(
        slide,
        "Flow-first policy",
        "GNN trunk is optimized for continuous flow, then asked to rank five actions.",
        Inches(0.58),
        y,
        Inches(2.65),
        Inches(1.12),
        CMU_RED,
    )
    add_horizontal_arrow(slide, Inches(3.38), Inches(2.1), Inches(0.65))
    add_callout(
        slide,
        "Retrofitted action head",
        "Action logits work best only after integrating the flow trajectory.",
        Inches(4.18),
        y,
        Inches(2.55),
        Inches(1.12),
        GOLD,
    )
    add_horizontal_arrow(slide, Inches(6.88), Inches(2.1), Inches(0.65))
    add_callout(
        slide,
        "Coordination gap",
        "Waiting and yielding decisions remain under-trained in hard crowds.",
        Inches(7.58),
        y,
        Inches(1.8),
        Inches(1.12),
        CMU_RED,
    )

    y2 = Inches(3.2)
    add_callout(
        slide,
        "Discrete imitation policy",
        "The shared trunk sees action cross-entropy throughout training.",
        Inches(0.58),
        y2,
        Inches(2.65),
        Inches(1.05),
        BLUE,
    )
    add_horizontal_arrow(slide, Inches(3.38), Inches(3.64), Inches(0.65), BLUE)
    add_callout(
        slide,
        "Coordination-native features",
        "Representations directly encode move, wait, and give-way choices.",
        Inches(4.18),
        y2,
        Inches(2.55),
        Inches(1.05),
        BLUE,
    )
    add_horizontal_arrow(slide, Inches(6.88), Inches(3.64), Inches(0.65), BLUE)
    add_callout(
        slide,
        "Dense-map strength",
        "Large random-map advantage suggests the objective matters.",
        Inches(7.58),
        y2,
        Inches(1.8),
        Inches(1.05),
        BLUE,
    )


def slide_experiment_ladder(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    add_title(
        slide,
        "Ablations: inference mismatch vs representation learning",
        "Each step asks whether the action head sees the same distribution at train and test.",
    )
    steps = [
        ("1", "Integrated action readout", "Query actions after the flow trajectory reaches the goal-side endpoint.", GREEN),
        ("2", "Endpoint action training", "Train discrete logits where inference reads them.", GREEN),
        ("3", "Discrete-primary hybrid", "Up-weight action CE to teach coordination earlier.", GOLD),
        ("4", "Balanced + larger trunk", "Keep the dense-map gain without sacrificing general maps.", BLUE),
    ]
    x0 = Inches(0.55)
    w = Inches(2.05)
    gap = Inches(0.26)
    for idx, title, body, color in steps:
        left = x0 + (w + gap) * (int(idx) - 1)
        add_label(slide, idx, left, Inches(1.52), Inches(0.36), Inches(0.36), color)
        add_callout(slide, title, body, left, Inches(1.96), w, Inches(1.45), color)
        if idx != "4":
            add_horizontal_arrow(slide, left + w + Inches(0.06), Inches(2.6), Inches(0.22), GRAY)

    add_placeholder(
        slide,
        "Plot placeholder:\nquick-vs-medium stability scatter",
        Inches(1.0),
        Inches(3.75),
        Inches(3.25),
        Inches(0.9),
    )
    add_text(
        slide,
        "Key lesson: quick evals can overstate wins; medium eval decides whether to spend full-eval compute.",
        Inches(4.65),
        Inches(3.92),
        Inches(4.0),
        Inches(0.42),
        size=13,
        bold=True,
        color=DARK,
    )


def slide_w2_results(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    add_title(
        slide,
        "Discrete pressure helps random-32, but hurts balance",
        "The w2 hybrid validates the hypothesis and exposes a generalization tradeoff.",
    )
    categories = ["Paris", "den312d", "den520d", "empty-48", "maze", "random-32", "random-64", "warehouse"]
    values = [82, 32, 100, 98, 40, 75, 62, 0]
    add_clustered_column_chart(
        slide,
        categories,
        [("Hybrid discrete-primary medium", values, CMU_RED)],
        Inches(0.45),
        Inches(1.68),
        Inches(6.5),
        Inches(2.9),
        legend=False,
    )
    add_custom_legend(
        slide,
        [("Hybrid discrete-primary medium", CMU_RED)],
        Inches(0.66),
        Inches(1.42),
    )
    add_callout(
        slide,
        "Good signal",
        "random-32 rises to 75%, well above the 50% flow-vector baseline.",
        Inches(7.08),
        Inches(1.62),
        Inches(2.28),
        Inches(0.92),
        GREEN,
    )
    add_callout(
        slide,
        "Bad signal",
        "General maps regress, especially Paris, maze, random-64, and warehouse.",
        Inches(7.08),
        Inches(2.78),
        Inches(2.28),
        Inches(1.05),
        CMU_RED,
    )
    add_text(
        slide,
        "Next experiment: lower the action weight and add capacity instead of pushing harder on CE.",
        Inches(7.08),
        Inches(4.08),
        Inches(2.28),
        Inches(0.52),
        size=12,
        color=GRAY,
    )


def slide_temperature_sweep(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    add_title(
        slide,
        "Inference temperature is a knob, not the fix",
        "tau = 0.5 recovers some general behavior but gives back the dense-map gain.",
    )
    add_clustered_column_chart(
        slide,
        ["Quick overall", "Medium overall", "Medium random-32", "Medium random-64"],
        [
            ("tau 0.3", [56.52, 60.00, 75, 62], CMU_RED),
            ("tau 0.5", [65.22, 61.08, 60, 66], BLUE),
        ],
        Inches(0.55),
        Inches(1.72),
        Inches(6.05),
        Inches(2.95),
        legend=False,
    )
    add_custom_legend(
        slide,
        [("tau 0.3", CMU_RED), ("tau 0.5", BLUE)],
        Inches(0.72),
        Inches(1.43),
        columns=2,
    )
    add_callout(
        slide,
        "Interpretation",
        "Temperature trades decisiveness for recovery; features must come from training.",
        Inches(6.88),
        Inches(1.72),
        Inches(2.35),
        Inches(1.12),
        BLUE,
    )
    add_placeholder(
        slide,
        "Plot placeholder:\ntau sweep by map family",
        Inches(6.88),
        Inches(3.16),
        Inches(2.35),
        Inches(1.28),
    )


def slide_next_decisions(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    add_title(
        slide,
        "Next decision: balance the loss, then scale if capacity helps",
        "Promote only models that improve medium performance without losing random-32.",
    )
    add_callout(
        slide,
        "Evaluate now",
        "Balanced hybrid, 256x3\nBalanced hybrid, 512x3",
        Inches(0.7),
        Inches(1.65),
        Inches(2.6),
        Inches(1.18),
        BLUE,
    )
    add_callout(
        slide,
        "Promote to full eval",
        "Medium overall >65%\nrandom-32 at least 70%",
        Inches(3.7),
        Inches(1.65),
        Inches(2.6),
        Inches(1.18),
        GREEN,
    )
    add_callout(
        slide,
        "Hold off",
        "Overall below 63%\nor random-32 falls near 50%",
        Inches(6.7),
        Inches(1.65),
        Inches(2.6),
        Inches(1.18),
        CMU_RED,
    )
    add_placeholder(
        slide,
        "Plot placeholder:\nmedium per-map comparison for 256x3 vs 512x3",
        Inches(0.85),
        Inches(3.2),
        Inches(4.0),
        Inches(1.28),
    )
    add_placeholder(
        slide,
        "Diagram placeholder:\ncapacity path: 512x3 -> 512x6 / 1024x6",
        Inches(5.18),
        Inches(3.2),
        Inches(3.7),
        Inches(1.28),
    )


def build_gap_analysis_deck() -> Path:
    if not INPUT_PPTX.exists():
        raise FileNotFoundError(f"Input presentation not found: {INPUT_PPTX}")

    DECK_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(INPUT_PPTX, OUTPUT_PPTX)

    prs = Presentation(str(OUTPUT_PPTX))
    slide_gap_headline(prs)
    slide_dense_coordination_gap(prs)
    slide_training_objective(prs)
    slide_experiment_ladder(prs)
    slide_w2_results(prs)
    slide_temperature_sweep(prs)
    slide_next_decisions(prs)
    prs.save(str(OUTPUT_PPTX))
    return OUTPUT_PPTX


if __name__ == "__main__":
    out_path = build_gap_analysis_deck()
    print(f"Wrote updated presentation: {out_path}")
